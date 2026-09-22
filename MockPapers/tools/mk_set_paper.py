"""Build one-page-per-set A4 Word answer sheets for the RT-3 mock sets.

    python tools/mk_set_paper.py 6            ->  RT3-Set6-Paper.docx
    python tools/mk_set_paper.py 6 7 8 9 10   ->  RT3-Sets6-10-Paper.docx (one page each)
    python tools/mk_set_paper.py CAT5 1 2 3   ->  CAT5-Sets1-3-Paper.docx

First arg may name the mock-set file's prefix (RT3, CAT4, CAT5...); defaults to RT3.

Columns: # | Physics code | Answer | Chemistry code | Answer | Maths code | Answer
Questions are numbered 1-25 straight through (MCQ + numerical merged, no part labels).
"""
import re, sys, os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

args = sys.argv[1:]
PREFIX = args.pop(0) if args and not args[0].isdigit() else 'RT3'
SETS = args or ['6']
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', '%s-Mock-Sets.html' % PREFIX)
OUT = os.path.join(HERE, '..', ('%s-Set%s-Paper.docx' % (PREFIX, SETS[0])) if len(SETS) == 1
                   else ('%s-Sets%s-%s-Paper.docx' % (PREFIX, SETS[0], SETS[-1])))

src = open(SRC, encoding='utf-8').read()
EXAM = re.search(r'<h1>(.*?)</h1>', src).group(1).replace(' Mock Sets', '')   # "ICAD CAT-5"

def codes_for(n):
    """{subject: [25 codes]} for set n, in book order (MCQ then numerical)."""
    i = src.find('id="s%s"' % n)
    j = src.find('id="s%d"' % (int(n) + 1))
    assert i > 0, 'set %s not found' % n
    block = src[i:j if j > i else len(src)].split('<details')[0]   # <details> = answer key
    subs = {sub[:sub.find('<')]: re.findall(r'class="code">([^<]+)<', sub)
            for sub in re.split(r'<h3>', block)[1:]}
    assert len(subs) == 3 and all(len(v) == 25 for v in subs.values()), \
        (n, {k: len(v) for k, v in subs.items()})
    return subs

# --- page --------------------------------------------------------------------
W_N, W_C, W_A = Mm(7), Mm(21), Mm(41)                      # #, code, answer
WIDTHS = [W_N, W_C, W_A, W_C, W_A, W_C, W_A]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.left_margin = sec.right_margin = Mm(8)
sec.bottom_margin = Mm(6)

def para(p, text, size, bold=False, mono=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size, r.font.bold = Pt(size), bold
    r.font.name = 'Consolas' if mono else 'Calibri'
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p

def shade(cell, fill):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(el)

def sheet(n, first):
    subs = codes_for(n)
    order = list(subs)

    title = doc.add_paragraph()
    if not first:
        title.add_run().add_break(WD_BREAK.PAGE)
    para(title, '%s  \u00b7  Mock Set %s \u2014 Answer Sheet' % (EXAM, n), 13, bold=True,
         align=WD_ALIGN_PARAGRAPH.LEFT).paragraph_format.space_after = Pt(3)
    meta = doc.add_paragraph()
    para(meta, 'Name: _______________________     Date: ____________     '
               'Time: 3 h     Score: _________', 9,
         align=WD_ALIGN_PARAGRAPH.LEFT, color='444444').paragraph_format.space_after = Pt(6)

    t = doc.add_table(rows=2, cols=7)
    t.style = 'Table Grid'
    t.autofit = False
    hdr, sub2 = t.rows[0], t.rows[1]
    hdr.cells[0].merge(sub2.cells[0])
    para(hdr.cells[0].paragraphs[0], '#', 9, bold=True)
    shade(hdr.cells[0], 'E0E0E0')
    for k, name in enumerate(order):
        c = hdr.cells[1 + 2 * k].merge(hdr.cells[2 + 2 * k])
        para(c.paragraphs[0], name, 10, bold=True)
        shade(c, 'D9D9D9')
        for off, lab in ((1, 'Q code'), (2, 'Answer')):
            c = sub2.cells[off + 2 * k]
            para(c.paragraphs[0], lab, 9, bold=True)
            shade(c, 'EFEFEF')

    for r in range(25):
        row = t.add_row()
        row.height = Mm(9.4)
        para(row.cells[0].paragraphs[0], str(r + 1), 7, color='555555')
        for k, name in enumerate(order):
            para(row.cells[1 + 2 * k].paragraphs[0], subs[name][r], 8.5, mono=True)

    for row in t.rows:
        for c, w in zip(row.cells, WIDTHS):
            c.width = w
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

for idx, n in enumerate(SETS):
    sheet(n, idx == 0)

doc.save(OUT)
print(OUT, len(SETS), 'sheet(s)')
